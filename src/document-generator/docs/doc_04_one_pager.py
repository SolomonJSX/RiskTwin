from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from styles import (
    apply_gost_page_setup, add_gost_paragraph, add_gost_heading,
    set_academic_table_borders, format_cell
)

def build_one_pager_document(output_path: str):
    doc = Document()
    apply_gost_page_setup(doc)

    # Title
    add_gost_heading(doc, "ONE-PAGER ПРОЕКТА RISKTWIN", space_before=0, space_after=3)
    add_gost_paragraph(
        doc,
        "Превентивная платформа управления кредитными рисками и реструктуризации на базе Neuro-Symbolic AI. "
        "Автор: Сулейменов Ж.С. (КарТУ им. Абылкаса Сагинова). Трек: AI for Finance (Конкурс АБРК & ГКБ 2026).",
        bold=False, space_after=6
    )

    # Таблица 1: Базовые характеристики проекта
    tbl_grid = doc.add_table(rows=4, cols=2)
    tbl_grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_grid.autofit = False
    set_academic_table_borders(tbl_grid)
    w_col = [Cm(4.5), Cm(12.5)]
    for row in tbl_grid.rows:
        row.cells[0].width = w_col[0]
        row.cells[1].width = w_col[1]

    blocks = [
        ("1. Проблема и контекст:", "Рост NPL в РК (8.5 млн заемщиков). Традиционный скоринг опаздывает на 30–90 дней, замораживая провизии банков по МСФО 9 (Stage 3)."),
        ("2. Целевая аудитория:", "Банки второго уровня РК (Bereke, Halyk, Kaspi), микрофинансовые организации (МФО), Государственное Кредитное Бюро (ГКБ), АРРФР."),
        ("3. Суть решения:", "Dual-Tower нейросеть (финансы + психометрия) + Stalker Engine (ранний стресс) + автоматический солвер 3 превентивных планов помощи в ДБО."),
        ("4. Ценность для банков:", "Удержание ссуд в Stage 2 (экономия капитала до 70%), рост возвратности долга на 35–40%, урегулирование без коллекторов и судов.")
    ]

    for idx, (lbl, val) in enumerate(blocks):
        format_cell(tbl_grid.cell(idx, 0), lbl, bold=True, size=Pt(12))
        format_cell(tbl_grid.cell(idx, 1), val, bold=False, size=Pt(12))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Раздел сравнения подходов
    add_gost_heading(doc, "Сравнительный анализ традиционного скоринга и платформы RiskTwin", space_before=4, space_after=3)

    tbl_comp = doc.add_table(rows=4, cols=3)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_comp.autofit = False
    set_academic_table_borders(tbl_comp)
    w_cells = [Cm(4.5), Cm(6.25), Cm(6.25)]
    for row in tbl_comp.rows:
        for j, w in enumerate(w_cells):
            row.cells[j].width = w

    comp_data = [
        ("Критерий сравнения", "Традиционный банковский скоринг", "Платформа RiskTwin"),
        ("Горизонт реагирования", "Постфактум (просрочка 30+ и 90+ DPD)", "Превентивно за 30 дней до дефолта"),
        ("Анализируемые сигналы", "Выписки счетов и история в ГКБ", "Поведение в ДБО, спад трат, МФО, стресс"),
        ("Механизм помощи", "Штрафы, звонки коллекторов, судебный иск", "Автоматический выбор 3 опций в приложении")
    ]

    for i, row in enumerate(comp_data):
        for j, val in enumerate(row):
            c = tbl_comp.cell(i, j)
            format_cell(c, val, bold=(i == 0 or j == 0), size=Pt(12))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Раздел метрик и стека
    add_gost_heading(doc, "Ключевые технические параметры и готовность (TRL 6–7)", space_before=4, space_after=2)
    add_gost_paragraph(
        doc,
        "Точность классификации раннего риска: AUC-ROC = 0.887 (Home Credit Default Risk). Точность прогноза кассового "
        "разрыва: MAE = ±4.2 дня. Стек разработки: .NET Core 8 WebAPI, gRPC, Next.js 14, Apache Kafka, PyTorch, LightGBM, "
        "TreeSHAP. Полная интеграционная готовность к пилотированию в банковской инфраструктуре Казахстана.",
        space_after=0
    )

    doc.save(output_path)
    print(f"[OK] Generated: {output_path}")

if __name__ == "__main__":
    build_one_pager_document("test_one_pager.docx")
