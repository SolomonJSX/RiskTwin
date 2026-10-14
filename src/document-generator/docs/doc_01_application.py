from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from styles import (
    apply_gost_page_setup, add_gost_paragraph, add_gost_heading,
    set_academic_table_borders, format_cell, FONT_NAME
)

def build_application_document(output_path: str):
    doc = Document()
    apply_gost_page_setup(doc)

    # 1. Шапка в правом углу (без абзацного отступа)
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.paragraph_format.line_spacing = 1.0
    p_head.paragraph_format.first_line_indent = Cm(0)
    p_head.paragraph_format.space_before = Pt(0)
    p_head.paragraph_format.space_after = Pt(18)

    r1 = p_head.add_run("В Организационный комитет\n")
    r1.font.bold = True
    r2 = p_head.add_run(
        "Республиканского конкурса студенческих проектов\n"
        "и исследований в финансовой сфере (АБРК & ГКБ)\n"
        "г. Алматы, ул. Панфилова, 98"
    )
    for r in (r1, r2):
        r.font.name = FONT_NAME
        r.font.size = Pt(12)

    # 2. Название документа (заголовок по ширине с отступом 0.75 см по требованию)
    add_gost_heading(doc, "ЗАЯВКА НА УЧАСТИЕ В РЕСПУБЛИКАНСКОМ КОНКУРСЕ", space_before=6, space_after=12)

    # 3. Вводный текст
    add_gost_paragraph(
        doc,
        "Карагандинский технический университет имени Абылкаса Сагинова (КарТУ) направляет "
        "студенческий проект-победитель внутривузовского отборочного этапа для участия во втором "
        "(финальном) туре Республиканского конкурса студенческих проектов и исследований в области "
        "искусственного интеллекта, финтех-разработок и цифровых решений в финансово-банковской сфере "
        "(г. Астана, Назарбаев Университет, 13 ноября 2026 года).",
        space_after=8
    )

    # 4. Таблица сведений (академическая сетка 0.5 пт, Times New Roman 12 пт)
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_academic_table_borders(table)

    col_widths = [Cm(5.0), Cm(12.0)]
    for row in table.rows:
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]

    data = [
        ("Организация образования (ОВПО):", "Карагандинский технический университет имени Абылкаса Сагинова (КарТУ)"),
        ("Полный адрес учебного заведения:", "100027, Республика Казахстан, Карагандинская область, г. Караганда, проспект Нурсултана Назарбаева, 56"),
        ("Тематический трек конкурса:", "AI for Finance: скоринг, прогнозирование рисков, LLM-сервисы, цифровая реструктуризация"),
        ("Название проекта:", "RiskTwin: Цифровой двойник кредитного риска и превентивный оптимизатор реструктуризации"),
        ("Автор разработки (участник):", "Сулейменов Жансерик Саматович, студент-исследователь КарТУ"),
        ("Форма представления проекта:", "Действующий программный MVP-комплекс, аналитический доклад, презентация, репозиторий"),
        ("Контактные данные заявителя:", "Тел.: +7 (705) 327-02-39 • E-mail: suleimenov.risk@gmail.com")
    ]

    for i, (label, val) in enumerate(data):
        format_cell(table.cell(i, 0), label, bold=True, size=Pt(12))
        format_cell(table.cell(i, 1), val, bold=False, size=Pt(12))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. Заключение вуза
    add_gost_heading(doc, "Заключение конкурсной комиссии ОВПО:", space_before=6, space_after=6)
    add_gost_paragraph(
        doc,
        "Представленная работа полностью соответствует критериям актуальности, научной новизны "
        "и практической применимости в банковском секторе Республики Казахстан, а также отвечает "
        "регуляторным требованиям АРРФР и стандартам МСФО 9. Проект рекомендован комиссией КарТУ "
        "к очной защите перед экспертным жюри на площадке Назарбаев Университета в г. Астана.",
        space_after=18
    )

    # 6. Блок подписей
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.line_spacing = 1.15
    p_sign.paragraph_format.first_line_indent = Cm(0)
    p_sign.paragraph_format.space_before = Pt(24)

    r_s = p_sign.add_run(
        "Проректор по научной работе КарТУ:      _________________ / _________________ /\n\n"
        "Ответственный координатор ОВПО:          _________________ / _________________ /\n\n"
        "Автор проекта (студент-исследователь):     _________________ / Сулейменов Ж.С. /\n\n"
        "Дата: «___» _____________ 2026 г.         М.П."
    )
    r_s.font.name = FONT_NAME
    r_s.font.size = Pt(12)

    doc.save(output_path)
    print(f"[OK] Generated: {output_path}")

if __name__ == "__main__":
    build_application_document("test_app.docx")
