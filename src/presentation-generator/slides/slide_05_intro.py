import os
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from theme import SLIDE_WIDTH, SLIDE_HEIGHT, FONT_MAIN

# Dark Theme Palette
COLOR_PURE_BLACK = RGBColor(0, 0, 0)
COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
COLOR_TEXT_MUTED_LIGHT = RGBColor(148, 163, 184) # slate-400
COLOR_EMERALD_GLOW = RGBColor(52, 211, 153)     # emerald-400
BORDER_DARK = RGBColor(39, 39, 42)              # zinc-800
BG_TAG_DARK = RGBColor(24, 24, 27)              # zinc-900

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Сплошной черный фон
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PURE_BLACK

    # 2. Аккуратный темный бейдж-тэг
    tag_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.0), Inches(0.85), Inches(4.8), Inches(0.38)
    )
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = BG_TAG_DARK
    tag_box.line.color.rgb = RGBColor(16, 185, 129)
    tag_box.line.width = Pt(1.2)
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = False
    tf_tag.margin_top = tf_tag.margin_bottom = Inches(0)
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "РАЗДЕЛ II • ТЕХНОЛОГИЧЕСКИЙ ФУНДАМЕНТ"
    p_tag.font.name = FONT_MAIN
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_EMERALD_GLOW
    p_tag.alignment = PP_ALIGN.CENTER

    # 3. ОГРОМНЫЙ ЗАГОЛОВОК ЗАГЛАВНЫМИ БУКВАМИ
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.38), Inches(11.333), Inches(1.4))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0

    p_title = tf_title.paragraphs[0]
    p_title.text = "ОЗНАКОМЛЕНИЕ С МОДЕЛЬЮ:"
    p_title.font.name = FONT_MAIN
    p_title.font.size = Pt(38)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE

    p_title2 = tf_title.add_paragraph()
    p_title2.text = "КАК СОЗДАВАЛСЯ И РАБОТАЕТ RISKTWIN"
    p_title2.font.name = FONT_MAIN
    p_title2.font.size = Pt(38)
    p_title2.font.bold = True
    p_title2.font.color.rgb = COLOR_EMERALD_GLOW

    # 4. Поясняющий подзаголовок (указывает, что следующие слайды — про создание модели)
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(11.333), Inches(0.75))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0

    p_sub = tf_sub.paragraphs[0]
    p_sub.text = (
        "В следующих слайдах мы подробно разберем внутреннюю архитектуру создания модели:\n"
        "от реляционной базы Home Credit (307K анкет) и поведенческой телеметрии — "
        "до двухбашенной нейросети PyTorch и математического солвера реструктуризации."
    )
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED_LIGHT

    # 5. ИЗОБРАЖЕНИЕ: Роза лежит на столе (атмосферный кинематографичный арт)
    rose_path = Path(__file__).resolve().parent.parent / "assets" / "rose_on_table.png"
    if rose_path.exists():
        # Размещаем по центру в нижней части слайда, гармонично растворяясь в черном фоне
        slide.shapes.add_picture(
            str(rose_path),
            Inches(2.5), Inches(3.75), Inches(8.333), Inches(3.35)
        )

    # 6. Минималистичный темный футер
    footer_box = slide.shapes.add_textbox(Inches(1.0), Inches(7.05), Inches(11.333), Inches(0.3))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "RiskTwin • Раздел: Создание и анатомия модели | Слайд 5 из 13"
    p_f.font.name = FONT_MAIN
    p_f.font.size = Pt(8.5)
    p_f.font.color.rgb = RGBColor(100, 116, 139)

    return slide
