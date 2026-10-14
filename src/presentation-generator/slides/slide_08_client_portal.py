from pptx.util import Inches, Pt
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_image_placeholder,
    add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    add_header(
        slide,
        tag_text="Клиентский опыт (ДБО)",
        title_text="Бесшовный клиентский путь: реструктуризация в 1 клик за 30 секунд",
        subtitle_text="Вместо психологического давления и звонков коллекторов — персонализированная помощь в приложении"
    )

    # Left Column: The 4 Client Steps with icon badges
    steps = [
        ("bell", "1. Деликатный Push в мобильном банке",
         "Банк не угрожает штрафами, а проявляет заботу: «Мы заметили временное изменение нагрузки. Для вас подготовлена индивидуальная программа поддержки»."),
        ("smartphone", "2. Интуитивная витрина 3 решений",
         "Заемщик заходит в приложение и видит 3 понятных карточки: точный новый платеж (₸89 500 вместо ₸145 000), новый срок, переплату и экономию."),
        ("key", "3. Смарт-подписание по SMS-OTP (код 7777)",
         "Юридически значимое подписание дополнительного соглашения прямо в мобильном банке за 10 секунд без визита в отделение."),
        ("check-circle", "4. Мгновенная активация и сохранение истории",
         "График в ДБО обновляется автоматически, статус кредита остается «зеленым» (Stage 1), кредитная история в ГКБ не портится.")
    ]

    for i, (icon_name, title, desc) in enumerate(steps):
        top = Inches(1.8 + i * 1.25)
        add_card(slide, Inches(0.8), top, Inches(6.0), Inches(1.18))
        
        # Icon badge
        add_icon_badge(
            slide, icon_name,
            Inches(1.0), top + Inches(0.18),
            badge_size=Inches(0.55), icon_size=Inches(0.32),
            bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
        )

        box = slide.shapes.add_textbox(Inches(1.7), top + Inches(0.08), Inches(4.9), Inches(1.02))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(2)

    # Right Column: Visual Frame for Client Offers View
    add_image_placeholder(
        slide,
        Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.9),
        "Экран «Витрина решений для клиента в ДБО»\n(3 варианта графика, калькулятор платежа\nи окно SMS-подписания по коду 7777)"
    )

    add_footer(slide, current_slide=8)
    return slide
