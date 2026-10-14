from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_image_placeholder,
    add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    COLOR_ROSE, COLOR_AMBER, BG_CARD, BG_MINT, BG_ROSE_LIGHT, BG_AMBER_LIGHT,
    BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    add_header(
        slide,
        tag_text="Сквозной кейс",
        title_text="Сквозной кейс: от скрытого стресса до мгновенного решения",
        subtitle_text="Как RiskTwin спасает надежного заемщика Армана (IT-Сеньор, Алматы) от сползания в дефолт"
    )

    # Left Column: Step-by-Step Flow with icon badges
    steps = [
        ("user", "Шаг 1: Исходное состояние (Зеленая зона)",
         "Доход: ₸1 200 000 | Долг: ₸4 200 000 | Платеж: ₸145 000. Базовый скоринг 1.3% (Класс AAA). Банк считает клиента абсолютно надежным.",
         BG_MINT, COLOR_PRIMARY_DARK, "#059669"),
        ("alert-triangle", "Шаг 2: Внезапный шок (Стресс-тест)",
         "Задержка расчета от зарубежного заказчика на 14 дней + непредвиденные медицинские расходы. Традиционный банк слеп (срок платежа еще не наступил).",
         BG_ROSE_LIGHT, COLOR_ROSE, "#E11D48"),
        ("bell", "Шаг 3: Раннее срабатывание RiskTwin",
         "Нейросеть фиксирует скрытый стресс по ночным входам в ДБО и падению медианного баланса. Модель вычисляет риск кассового разрыва за 30 дней.",
         BG_AMBER_LIGHT, COLOR_AMBER, "#D97706"),
        ("shield-check", "Шаг 4: Алгоритмическое спасение (OR-Tools)",
         "Решатель за 12ms формирует соглашение: продление срока на 18 мес., платеж снижается со ₸145 000 до ₸89 500. Кредит спасен, провизии не создаются.",
         BG_MINT, COLOR_PRIMARY_DARK, "#059669")
    ]

    for i, (icon_name, s_title, s_desc, bg_c, text_c, hex_c) in enumerate(steps):
        top = Inches(1.8 + i * 1.25)
        add_card(slide, Inches(0.8), top, Inches(6.0), Inches(1.18), bg_color=bg_c)
        
        # Icon badge
        add_icon_badge(
            slide, icon_name,
            Inches(1.0), top + Inches(0.18),
            badge_size=Inches(0.55), icon_size=Inches(0.32),
            bg_color=bg_c, border_color=text_c, color_hex=hex_c
        )

        box = slide.shapes.add_textbox(Inches(1.7), top + Inches(0.08), Inches(4.9), Inches(1.02))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = s_title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = text_c

        p2 = tf.add_paragraph()
        p2.text = s_desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = COLOR_TEXT_DARK
        p2.space_before = Pt(2)

    # Right Column: Visual Frame for Stress Simulator & Chart
    from pathlib import Path
    img_path = Path(__file__).resolve().parent.parent / "assets" / "ui_stress_simulator.png"
    add_card(slide, Inches(7.05), Inches(1.8), Inches(5.45), Inches(4.95), bg_color=RGBColor(248, 250, 252), border_color=BORDER_LIGHT)
    if img_path.exists():
        slide.shapes.add_picture(
            str(img_path),
            Inches(7.10), Inches(1.85), Inches(5.35), Inches(4.85)
        )
    else:
        add_image_placeholder(
            slide,
            Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.9),
            "Экран «Стресс-тестирование & Прогноз ликвидности»"
        )

    add_footer(slide, current_slide=16, total_slides=20)
    return slide
