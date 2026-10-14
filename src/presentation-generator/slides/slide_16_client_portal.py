import os
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    # 1. Заголовок слайда с бейджем
    add_header(
        slide,
        tag_text="РАЗДЕЛ IV • ЖИВОЙ ИНТЕРФЕЙС ПЛАТФОРМЫ",
        title_text="Бесшовный клиентский путь (ДБО): реструктуризация в 1 клик за 30 секунд",
        subtitle_text="Сквозной процесс: от раннего алерта риск-офицера до подписания нового договора по SMS в мобильном банке",
        title_size=Pt(18.5)
    )

    assets_dir = Path(__file__).resolve().parent.parent / "assets"
    img1_path = assets_dir / "ui_portal_step1.png"
    img2_path = assets_dir / "ui_portal_step2.png"
    img3_path = assets_dir / "ui_portal_step3.png"
    img4_path = assets_dir / "ui_portal_step4.png"

    col_width = Inches(2.88)
    gap = Inches(0.18)
    start_x = Inches(0.60)

    # Данные по 4 шагам клиентского пути
    steps_data = [
        {
            "img": img1_path,
            "badge_icon": "alert-circle",
            "badge_color": "#DC2626",
            "badge_bg": RGBColor(254, 242, 242),
            "badge_border": RGBColor(254, 202, 202),
            "card_bg": RGBColor(254, 242, 242),
            "card_border": RGBColor(254, 202, 202),
            "title_color": RGBColor(185, 28, 28),
            "title": "Шаг 1: Алерт EWS & Ожидание",
            "desc": (
                "• EWS перехватывает скрытый стресс Армана: задержка ЗП +14 дней, падение подушки -35%.\n"
                "• Модель OR-Tools синтезирует 3 щадящих графика (₸78k, ₸108k, ₸89k).\n"
                "• Смартфон клиента заблокирован: ожидает санкции риск-офицера банка."
            )
        },
        {
            "img": img2_path,
            "badge_icon": "bell",
            "badge_color": "#2563EB",
            "badge_bg": RGBColor(239, 246, 255),
            "badge_border": RGBColor(191, 219, 254),
            "card_bg": RGBColor(239, 246, 255),
            "card_border": RGBColor(191, 219, 254),
            "title_color": RGBColor(29, 78, 216),
            "title": "Шаг 2: Санкция Kafka & Push",
            "desc": (
                "• Риск-офицер одобряет план (Human-in-the-Loop) в 1 клик.\n"
                "• Событие публикуется в Kafka топик risk.restructuring.approved.v1.\n"
                "• На экране смартфона клиента мгновенно всплывает деликатный Push банка о персональной поддержке."
            )
        },
        {
            "img": img3_path,
            "badge_icon": "smartphone",
            "badge_color": "#059669",
            "badge_bg": BG_MINT,
            "badge_border": BORDER_EMERALD,
            "card_bg": BG_MINT,
            "card_border": BORDER_EMERALD,
            "title_color": COLOR_PRIMARY_DARK,
            "title": "Шаг 3: Мобильная витрина 3 опций",
            "desc": (
                "• Клиент открывает мобильный банк и видит прозрачную витрину 3 решений:\n"
                "  1) Снижение платежа до ₸78 400\n"
                "  2) Каникулы 3 мес с платежом ₸0 (далее ₸108 200)\n"
                "  3) Оптимальный баланс ₸89 500\n"
                "• Клиент выбирает вариант «Каникулы»."
            )
        },
        {
            "img": img4_path,
            "badge_icon": "check-circle",
            "badge_color": "#15803D",
            "badge_bg": RGBColor(240, 253, 244),
            "badge_border": RGBColor(187, 247, 208),
            "card_bg": RGBColor(240, 253, 244),
            "card_border": RGBColor(187, 247, 208),
            "title_color": RGBColor(21, 128, 61),
            "title": "Шаг 4: SMS-OTP и спасение",
            "desc": (
                "• Всплывает модальное окно смарт-подписания электронного допсоглашения.\n"
                "• Заемщик вводит SMS-код 7777 и подтверждает договор за 5 секунд.\n"
                "• График обновлен: 0 дней просрочки, статус Stage 1 сохранен, кредитная история в ГКБ чиста."
            )
        }
    ]

    for i, step in enumerate(steps_data):
        col_x = start_x + i * (col_width + gap)

        # 1. Верхняя карточка со скриншотом
        add_card(slide, col_x, Inches(1.56), col_width, Inches(2.68), bg_color=RGBColor(248, 250, 252), border_color=BORDER_LIGHT)
        if step["img"].exists():
            slide.shapes.add_picture(
                str(step["img"]),
                col_x + Inches(0.04), Inches(1.60), col_width - Inches(0.08), Inches(2.60)
            )

        # 2. Нижняя карточка с описанием
        add_card(slide, col_x, Inches(4.30), col_width, Inches(2.14), bg_color=step["card_bg"], border_color=step["card_border"])

        add_icon_badge(
            slide, step["badge_icon"],
            col_x + Inches(0.10), Inches(4.38),
            badge_size=Inches(0.38), icon_size=Inches(0.24),
            bg_color=step["badge_bg"], border_color=step["badge_border"], color_hex=step["badge_color"]
        )

        box = slide.shapes.add_textbox(col_x + Inches(0.52), Inches(4.34), col_width - Inches(0.58), Inches(2.04))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_h = tf.paragraphs[0]
        p_h.text = step["title"]
        p_h.font.name = FONT_MAIN
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = step["title_color"]

        p_desc = tf.add_paragraph()
        p_desc.text = step["desc"]
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(7.4)
        p_desc.font.color.rgb = COLOR_TEXT_DARK
        p_desc.space_before = Pt(2.0)

    # 3. Нижний статус-бар
    bot_card = add_card(slide, Inches(0.60), Inches(6.50), Inches(12.06), Inches(0.40), bg_color=RGBColor(241, 245, 249), border_color=BORDER_LIGHT)
    bot_box = slide.shapes.add_textbox(Inches(0.70), Inches(6.53), Inches(11.86), Inches(0.34))
    tf_bot = bot_box.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_right = tf_bot.margin_top = tf_bot.margin_bottom = 0
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "⚡ Итог внедрения: 0 звонков коллекторов • 0 визитов в отделение • 100% STP в ДБО • Подписание договора за 30 секунд"
    p_bot.font.name = FONT_MAIN
    p_bot.font.size = Pt(8.3)
    p_bot.font.bold = True
    p_bot.font.color.rgb = COLOR_PRIMARY_DARK

    add_footer(slide, current_slide=16, total_slides=20)
    return slide
