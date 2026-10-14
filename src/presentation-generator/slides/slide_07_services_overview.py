import os
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    # 1. Заголовок слайда с бейджем
    add_header(
        slide,
        tag_text="АРХИТЕКТУРНЫЕ МОДУЛИ • РАЗДЕЛ III",
        title_text="Связка ключевых сервисов: Psycho-Life & Stalker Engine",
        subtitle_text="Двухконтурная система: психометрическая детекция скрытого стресса и реактивный радар мониторинга 24/7",
        title_size=Pt(19.5)
    )

    # -------------------------------------------------------------------------
    # СТРОКА 1: Схема кластера обоих сервисов (Векторный рендер 1050 x 320)
    # -------------------------------------------------------------------------
    cluster_path = Path(__file__).resolve().parent.parent / "assets" / "services_cluster.png"
    if cluster_path.exists():
        slide.shapes.add_picture(
            str(cluster_path),
            Inches(0.8), Inches(1.72), Inches(11.733), Inches(2.42)
        )

    # -------------------------------------------------------------------------
    # СТРОКА 2: Два информационных блока о названиях и назначении сервисов
    # -------------------------------------------------------------------------

    # Блок 1: Psycho-Life Service (Слева, фиолетовый акцент)
    add_card(slide, Inches(0.8), Inches(4.24), Inches(5.76), Inches(1.05), bg_color=RGBColor(250, 245, 255), border_color=RGBColor(233, 213, 255))
    
    add_icon_badge(
        slide, "brain",
        Inches(0.95), Inches(4.38),
        badge_size=Inches(0.48), icon_size=Inches(0.28),
        bg_color=RGBColor(250, 245, 255), border_color=RGBColor(216, 180, 254), color_hex="#7E22CE"
    )

    box_s1 = slide.shapes.add_textbox(Inches(1.55), Inches(4.3), Inches(4.9), Inches(0.95))
    tf_s1 = box_s1.text_frame
    tf_s1.word_wrap = True
    tf_s1.margin_left = tf_s1.margin_right = tf_s1.margin_top = tf_s1.margin_bottom = 0

    p_s1h = tf_s1.paragraphs[0]
    p_s1h.text = "1. Psycho-Life Service (Психометрический анализатор)"
    p_s1h.font.name = FONT_MAIN
    p_s1h.font.size = Pt(10.5)
    p_s1h.font.bold = True
    p_s1h.font.color.rgb = RGBColor(126, 34, 206)

    p_s1b = tf_s1.add_paragraph()
    p_s1b.text = (
        "Вычисляет невидимый в классических отчетах стресс заемщика: ночные сессии в ДБО (3 AM), "
        "экранное время на графике выплат, скорость падения остатка 7d и микрозаймы в МФО."
    )
    p_s1b.font.name = FONT_MAIN
    p_s1b.font.size = Pt(8.4)
    p_s1b.font.color.rgb = COLOR_TEXT_DARK
    p_s1b.space_before = Pt(2)

    # Блок 2: Stalker Engine (Справа, красный акцент)
    add_card(slide, Inches(6.77), Inches(4.24), Inches(5.76), Inches(1.05), bg_color=RGBColor(254, 242, 242), border_color=RGBColor(254, 202, 202))

    add_icon_badge(
        slide, "shield-check",
        Inches(6.92), Inches(4.38),
        badge_size=Inches(0.48), icon_size=Inches(0.28),
        bg_color=RGBColor(254, 242, 242), border_color=RGBColor(252, 165, 165), color_hex="#DC2626"
    )

    box_s2 = slide.shapes.add_textbox(Inches(7.52), Inches(4.3), Inches(4.9), Inches(0.95))
    tf_s2 = box_s2.text_frame
    tf_s2.word_wrap = True
    tf_s2.margin_left = tf_s2.margin_right = tf_s2.margin_top = tf_s2.margin_bottom = 0

    p_s2h = tf_s2.paragraphs[0]
    p_s2h.text = "2. Stalker Engine (Потоковый радар мониторинга 24/7)"
    p_s2h.font.name = FONT_MAIN
    p_s2h.font.size = Pt(10.5)
    p_s2h.font.bold = True
    p_s2h.font.color.rgb = RGBColor(185, 28, 28)

    p_s2b = tf_s2.add_paragraph()
    p_s2b.text = (
        "Реактивный движок на .NET 8 и SignalR WebSockets. Непрерывно принимает телеметрию, "
        "ведет таймлайн событий клиента, поднимает алерт риск-офицеру и мгновенно триггерит солвер."
    )
    p_s2b.font.name = FONT_MAIN
    p_s2b.font.size = Pt(8.4)
    p_s2b.font.color.rgb = COLOR_TEXT_DARK
    p_s2b.space_before = Pt(2)

    # -------------------------------------------------------------------------
    # СТРОКА 3: Описание от своего лица (Авторская мысль, серый цвет, Italic)
    # -------------------------------------------------------------------------
    add_card(slide, Inches(0.8), Inches(5.38), Inches(11.733), Inches(1.52), bg_color=RGBColor(248, 250, 252), border_color=RGBColor(203, 213, 225))

    add_icon_badge(
        slide, "sparkles",
        Inches(0.98), Inches(5.55),
        badge_size=Inches(0.52), icon_size=Inches(0.32),
        bg_color=RGBColor(241, 245, 249), border_color=RGBColor(148, 163, 184), color_hex="#475569"
    )

    box_author = slide.shapes.add_textbox(Inches(1.65), Inches(5.46), Inches(10.7), Inches(1.36))
    tf_a = box_author.text_frame
    tf_a.word_wrap = True
    tf_a.margin_left = tf_a.margin_right = tf_a.margin_top = tf_a.margin_bottom = 0

    p_ah = tf_a.paragraphs[0]
    p_ah.text = "Авторское видение: почему программе потребовалось разделение на два модуля"
    p_ah.font.name = FONT_MAIN
    p_ah.font.size = Pt(10.5)
    p_ah.font.bold = True
    p_ah.font.color.rgb = RGBColor(51, 65, 85)

    p_ab = tf_a.add_paragraph()
    p_ab.text = (
        "«Чтобы кардинально улучшить работоспособность и скорость реакции моей программы, я осознал: "
        "нельзя опираться только на статичный скоринг, который пересчитывается раз в месяц. Кредитный риск живой — "
        "он зарождается задолго до просрочки, когда заемщик в панике не спит ночью и берет первые микрозаймы. "
        "Поэтому я разделил систему на два специализированных органа: Psycho-Life Service стал чувствительными «нервами», "
        "считывающими эмоционально-финансовый стресс в мобильном банке, а Stalker Engine — неусыпным «радаром», "
        "который 24/7 держит портфель на контроле и без секунды промедления передает эстафету оптимизатору реструктуризации»."
    )
    p_ab.font.name = FONT_MAIN
    p_ab.font.size = Pt(9.8)
    p_ab.font.italic = True
    p_ab.font.color.rgb = RGBColor(71, 85, 105)
    p_ab.space_before = Pt(3.5)

    add_footer(slide, current_slide=9, total_slides=14)
    return slide
