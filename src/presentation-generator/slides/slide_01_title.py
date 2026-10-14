from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, apply_image_background, add_footer, add_card,
    add_icon_badge, add_icon,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_WHITE, BG_MINT, BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # 1. Full-bleed Frosted Glass Banking Architecture Background
    bg_path = Path(__file__).resolve().parent.parent / "assets" / "slide2_blurred_bg.jpg"
    if bg_path.exists():
        apply_image_background(slide, str(bg_path))
    else:
        apply_slide_background(slide)

    # 2. Category Badge with Award Icon
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.0), Inches(0.95), Inches(5.8), Inches(0.42)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = BG_MINT
    badge.line.color.rgb = BORDER_EMERALD
    add_icon(slide, "award", Inches(1.15), Inches(1.03), size=Inches(0.24), color_hex="#047857")

    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "     РЕСПУБЛИКАНСКИЙ КОНКУРС АБРК & ГКБ 2026"
    p_b.font.name = FONT_MAIN
    p_b.font.size = Pt(9.5)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_PRIMARY_DARK
    p_b.alignment = PP_ALIGN.LEFT

    # 3. Main Title Text Box with Shield Icon Badge
    add_icon_badge(
        slide, "shield-check",
        Inches(1.0), Inches(1.6),
        badge_size=Inches(0.85), icon_size=Inches(0.52),
        bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
    )

    title_box = slide.shapes.add_textbox(Inches(2.05), Inches(1.5), Inches(10.2), Inches(2.1))
    tf = title_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "RiskTwin"
    p1.font.name = FONT_MAIN
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "Превентивная платформа управления кредитными рисками и реструктуризации"
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(21)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_DARK
    p2.space_before = Pt(4)

    # 4. Description / Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.72), Inches(11.3), Inches(0.8))
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = "Интеллектуальная система превентивного раннего выявления кассовых разрывов заемщиков, симуляции стресс-сценариев и автономного урегулирования задолженности по стандартам АРРФР и МСФО 9."
    p_s.font.name = FONT_MAIN
    p_s.font.size = Pt(13)
    p_s.font.color.rgb = COLOR_TEXT_MUTED

    # 5. Badges / Track Chips
    chip1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.62), Inches(4.5), Inches(0.38))
    chip1.fill.solid()
    chip1.fill.fore_color.rgb = BG_WHITE
    chip1.line.color.rgb = BORDER_LIGHT
    add_icon(slide, "cpu", Inches(1.12), Inches(4.7), size=Inches(0.2), color_hex="#059669")
    tf_c1 = chip1.text_frame
    p_c1 = tf_c1.paragraphs[0]
    p_c1.text = "      Трек: AI for Finance / Digital Banking & RegTech"
    p_c1.font.name = FONT_MAIN
    p_c1.font.size = Pt(9.5)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_PRIMARY_DARK

    chip2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.62), Inches(4.62), Inches(3.4), Inches(0.38))
    chip2.fill.solid()
    chip2.fill.fore_color.rgb = BG_WHITE
    chip2.line.color.rgb = BORDER_LIGHT
    add_icon(slide, "scale", Inches(5.74), Inches(4.7), size=Inches(0.2), color_hex="#059669")
    tf_c2 = chip2.text_frame
    p_c2 = tf_c2.paragraphs[0]
    p_c2.text = "      Стандарты: АРРФР РК & МСФО 9"
    p_c2.font.name = FONT_MAIN
    p_c2.font.size = Pt(9.5)
    p_c2.font.bold = True
    p_c2.font.color.rgb = COLOR_TEXT_DARK

    chip3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.14), Inches(4.62), Inches(3.19), Inches(0.38))
    chip3.fill.solid()
    chip3.fill.fore_color.rgb = BG_WHITE
    chip3.line.color.rgb = BORDER_LIGHT
    add_icon(slide, "layers", Inches(9.26), Inches(4.7), size=Inches(0.2), color_hex="#059669")
    tf_c3 = chip3.text_frame
    p_c3 = tf_c3.paragraphs[0]
    p_c3.text = "      Twin Core: Neuro-Symbolic AI"
    p_c3.font.name = FONT_MAIN
    p_c3.font.size = Pt(9.5)
    p_c3.font.bold = True
    p_c3.font.color.rgb = COLOR_TEXT_DARK

    # 6. Full-Width Clean Author Card
    add_card(slide, Inches(1.0), Inches(5.25), Inches(11.333), Inches(1.45), bg_color=BG_WHITE, border_color=BORDER_LIGHT)
    add_icon_badge(
        slide, "user-check",
        Inches(1.22), Inches(5.5),
        badge_size=Inches(0.85), icon_size=Inches(0.48),
        bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
    )

    meta_box = slide.shapes.add_textbox(Inches(2.25), Inches(5.35), Inches(9.8), Inches(1.25))
    tf_m = meta_box.text_frame
    tf_m.word_wrap = True
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "Автор разработки: Сулейменов Ж.С."
    p_m1.font.name = FONT_MAIN
    p_m1.font.size = Pt(13)
    p_m1.font.bold = True
    p_m1.font.color.rgb = COLOR_TEXT_DARK

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Карагандинский технический университет имени Абылкаса Сагинова (КарТУ)"
    p_m2.font.name = FONT_MAIN
    p_m2.font.size = Pt(11.5)
    p_m2.font.color.rgb = COLOR_TEXT_MUTED
    p_m2.space_before = Pt(3)

    p_m3 = tf_m.add_paragraph()
    p_m3.text = "Локация / Год: Караганда — Астана, 2026"
    p_m3.font.name = FONT_MAIN
    p_m3.font.size = Pt(10.5)
    p_m3.font.bold = True
    p_m3.font.color.rgb = COLOR_PRIMARY_DARK
    p_m3.space_before = Pt(2)

    add_footer(slide, current_slide=1)
    return slide
