from pptx.util import Inches, Pt
from theme import (
    apply_slide_background, add_header, add_footer, add_card,
    add_icon_badge, add_icon,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    add_header(
        slide,
        tag_text="План внедрения & Финал",
        title_text="Дорожная карта пилотирования и масштабирования (6 месяцев)",
        subtitle_text="План интеграции RiskTwin в казахстанский банковский сектор в партнерстве с АБРК и ГКБ"
    )

    phases = [
        ("database", "Этап 1: Месяцы 1–2",
         "Песочница & Интеграция",
         "• Подключение к API Государственного Кредитного Бюро (ГКБ)\n• Калибровка весов нейросети на обезличенном историческом портфеле банка-партнера\n• Прохождение комплаенс-проверки безопасности"),
        ("target", "Этап 2: Месяцы 3–4",
         "A/B пилот (10 000 заемщиков)",
         "• Запуск превентивного мониторинга на контрольной группе 10 000 кредитов\n• Оценка опережения дефолтов (метрики EWS за 30–45 дней)\n• Тестирование конверсии предложений реструктуризации в ДБО"),
        ("rocket", "Этап 3: Месяцы 5–6",
         "Масштабирование в Production",
         "• Полномасштабный релиз на весь розничный портфель банка\n• Бесшовное подписание соглашений через Open Banking API\n• Регуляторная отчетность в АРРФР со снижением резервов")
    ]

    for i, (icon_name, p_tag, p_title, p_desc) in enumerate(phases):
        left = Inches(0.8 + i * 4.0)
        card = add_card(slide, left, Inches(1.8), Inches(3.7), Inches(3.2))
        
        # Icon badge
        add_icon_badge(
            slide, icon_name,
            left + Inches(0.2), Inches(1.95),
            badge_size=Inches(0.6), icon_size=Inches(0.34),
            bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
        )

        box = slide.shapes.add_textbox(left + Inches(0.9), Inches(1.95), Inches(2.6), Inches(0.6))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = p_tag.upper()
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = p_title
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(11.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_DARK

        desc_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.65), Inches(3.3), Inches(2.2))
        tf_d = desc_box.text_frame
        tf_d.word_wrap = True
        p3 = tf_d.paragraphs[0]
        p3.text = p_desc
        p3.font.name = FONT_MAIN
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Contact & Value Banner with Trophy Icon Badge
    add_card(slide, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.6), bg_color=BG_MINT, border_color=BORDER_EMERALD)
    add_icon_badge(
        slide, "award",
        Inches(1.05), Inches(5.45),
        badge_size=Inches(0.85), icon_size=Inches(0.48),
        bg_color=BG_MINT, border_color=COLOR_PRIMARY_DARK, color_hex="#047857"
    )

    c_box = slide.shapes.add_textbox(Inches(2.05), Inches(5.28), Inches(10.2), Inches(1.4))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True

    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "RiskTwin: Сохраняем капитал банков Казахстана, защищая финансовое будущее заемщиков"
    p_c1.font.name = FONT_MAIN
    p_c1.font.size = Pt(14)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_PRIMARY_DARK

    p_c2 = tf_c.add_paragraph()
    p_c2.text = "Автор: Сулейменов Ж.С. | Карагандинский технический университет имени Абылкаса Сагинова (КарТУ) | Караганда — Астана, 2026"
    p_c2.font.name = FONT_MAIN
    p_c2.font.size = Pt(10)
    p_c2.font.color.rgb = COLOR_TEXT_DARK
    p_c2.space_before = Pt(3)

    p_c3 = tf_c.add_paragraph()
    p_c3.text = "Готовы к демонстрации действующего прототипа и ответам на вопросы жюри."
    p_c3.font.name = FONT_MAIN
    p_c3.font.size = Pt(10)
    p_c3.font.italic = True
    p_c3.font.color.rgb = COLOR_TEXT_MUTED
    p_c3.space_before = Pt(2)

    add_footer(slide, current_slide=11)
    return slide
