from pptx.util import Inches, Pt
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_image_placeholder,
    add_icon_badge, add_icon,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    add_header(
        slide,
        tag_text="AI & Deep Learning",
        title_text="Архитектура модели: PyTorch Full 360° Cross-Attention Net",
        subtitle_text="Двухбашенная нейросеть с механизмом внимания над статическими и динамическими данными"
    )

    # Left Column: Neural Network Components (3 cards with icon badges)
    components = [
        ("database", "Башня статических факторов (Static Tower)",
         "Обрабатывает 15 признаков: кредитную историю ГКБ/ПКБ, DTI, стаж, возраст, класс надежности, параметры одобренных и отклоненных займов. Эмбеддинг: d=128."),
        ("bar-chart", "Башня динамических признаков (Behavioral Tower)",
         "Обрабатывает временные ряды: частоту ночных входов в ДБО, снятие наличных в банкоматах, падение остатка 7d/30d, среднюю задержку оплат по рассрочкам."),
        ("brain", "Слой Cross-Attention & Выходной проектор",
         "Механизм внимания выявляет скрытые корреляции между финансовой устойчивостью и поведенческими аномалиями. Выдает калиброванный PD (вероятность дефолта).")
    ]

    for i, (icon_name, c_title, c_desc) in enumerate(components):
        top = Inches(1.8 + i * 1.35)
        add_card(slide, Inches(0.8), top, Inches(6.0), Inches(1.25))
        
        # Icon badge
        add_icon_badge(
            slide, icon_name,
            Inches(1.0), top + Inches(0.18),
            badge_size=Inches(0.55), icon_size=Inches(0.32),
            bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
        )

        box = slide.shapes.add_textbox(Inches(1.7), top + Inches(0.1), Inches(4.9), Inches(1.05))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = c_title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = c_desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(2)

    # Bottom Metric Bar with Zap Icon
    add_card(slide, Inches(0.8), Inches(5.9), Inches(6.0), Inches(0.8), bg_color=BG_MINT, border_color=BORDER_EMERALD)
    add_icon(slide, "zap", Inches(1.0), Inches(6.12), size=Inches(0.35), color_hex="#047857")
    m_box = slide.shapes.add_textbox(Inches(1.45), Inches(5.95), Inches(5.2), Inches(0.7))
    tf_m = m_box.text_frame
    p_m = tf_m.paragraphs[0]
    p_m.text = "Метрики модели: Инференс ~7.4 ms | ROC-AUC: 75.6% | Интерпретируемость (XAI: SHAP)"
    p_m.font.name = FONT_MAIN
    p_m.font.size = Pt(10)
    p_m.font.bold = True
    p_m.font.color.rgb = COLOR_PRIMARY_DARK

    # Right Column: Visual Frame for Attention Heatmap
    add_image_placeholder(
        slide,
        Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.9),
        "Визуализация архитектуры & Вектор внимания\n(Тепловая карта Deep Learning Explainer\nс весами событий во времени)"
    )

    add_footer(slide, current_slide=5)
    return slide
