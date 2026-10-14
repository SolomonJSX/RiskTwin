import os
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_MINT, BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    # 1. Заголовок слайда с бейджем
    add_header(
        slide,
        tag_text="ОФИЦИАЛЬНАЯ ПРЕЗЕНТАЦИЯ РЕШЕНИЯ • HERO REVEAL",
        title_text="Встречайте — RiskTwin: Цифровой двойник кредитного риска",
        subtitle_text="Идея, родившаяся из моих поисков и доводов: спасти заемщика и защитить капитал банков Казахстана",
        title_size=Pt(19.5)
    )

    # -------------------------------------------------------------------------
    # ЛЕВАЯ СТОРОНА: Крупный план логотипа RiskTwin (Ромб + Зигзаги Halyk-стиль)
    # -------------------------------------------------------------------------
    logo_path = Path(__file__).resolve().parent.parent / "assets" / "risktwin_logo_hero.png"
    if logo_path.exists():
        slide.shapes.add_picture(
            str(logo_path),
            Inches(0.8), Inches(1.75), Inches(5.1), Inches(5.1)
        )

    # -------------------------------------------------------------------------
    # ПРАВАЯ СТОРОНА: Манифест автора и 3 технологических столпа решения
    # -------------------------------------------------------------------------
    
    # 1. Верхняя карточка-манифест с цитатой автора (без упоминания Palantir)
    add_card(slide, Inches(6.2), Inches(1.75), Inches(6.3), Inches(1.42), bg_color=BG_MINT, border_color=BORDER_EMERALD)
    
    add_icon_badge(
        slide, "sparkles",
        Inches(6.38), Inches(1.92),
        badge_size=Inches(0.52), icon_size=Inches(0.32),
        bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#047857"
    )

    box_quote = slide.shapes.add_textbox(Inches(7.05), Inches(1.85), Inches(5.35), Inches(1.25))
    tf_q = box_quote.text_frame
    tf_q.word_wrap = True
    tf_q.margin_left = tf_q.margin_right = tf_q.margin_top = tf_q.margin_bottom = 0

    p_qh = tf_q.paragraphs[0]
    p_qh.text = "«Встречайте, RiskTwin — результат моих доводов и мыслей»"
    p_qh.font.name = FONT_MAIN
    p_qh.font.size = Pt(12)
    p_qh.font.bold = True
    p_qh.font.color.rgb = COLOR_PRIMARY_DARK

    p_qb = tf_q.add_paragraph()
    p_qb.text = (
        "Мы создали превентивную экосистему на стыке цифровых двойников, глубоких нейросетей и исследования операций. "
        "RiskTwin меняет парадигму: вместо запоздалого взыскания через 90 дней система превентивно "
        "оздоравливает заемщика за месяц до дефолта, сохраняя финансовую стабильность и капитал банка."
    )
    p_qb.font.name = FONT_MAIN
    p_qb.font.size = Pt(9.5)
    p_qb.font.italic = True
    p_qb.font.color.rgb = COLOR_TEXT_DARK
    p_qb.space_before = Pt(4)

    # 2. Три ключевых столпа решения (с подробным раскрытием: Что такое цифровой двойник)
    pillars = [
        ("layers", "1. Динамический Цифровой Двойник (360° Ontology)",
         "Непрерывно агрегирует транзакции, P2P-переводы, чеки и поведение в приложении ДБО в единую живую модель финансового состояния заемщика, позволяя прогнозировать кассовый разрыв заемщика в реальном времени."),
        ("zap", "2. Превентивное предупреждение за 30–60 дней (EWS)",
         "Двухбашенная нейросеть с кросс-вниманием (PyTorch) вычисляет скрытый стресс с точностью ROC-AUC 0.941 задолго до 90 дней, когда заемщик еще платит по графику."),
        ("calculator", "3. Автоматическая реструктуризация (Prescriptive AI)",
         "Математический солвер (MILP) за 12 мс синтезирует щадящий график со 100% соблюдением нормативов АРРФР (КДН ≤ 50%), доставляя оффер прямо в мобильный банк.")
    ]

    for i, (icon_name, p_title, p_desc) in enumerate(pillars):
        top = Inches(3.32 + i * 1.18)
        add_card(slide, Inches(6.2), top, Inches(6.3), Inches(1.08), bg_color=RGBColor(248, 250, 252), border_color=BORDER_LIGHT)

        add_icon_badge(
            slide, icon_name,
            Inches(6.38), top + Inches(0.18),
            badge_size=Inches(0.55), icon_size=Inches(0.32),
            bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#047857"
        )

        box_p = slide.shapes.add_textbox(Inches(7.1), top + Inches(0.12), Inches(5.3), Inches(0.88))
        tf_p = box_p.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0

        p1 = tf_p.paragraphs[0]
        p1.text = p_title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf_p.add_paragraph()
        p2.text = p_desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(8.6)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(2.5)

    add_footer(slide, current_slide=4)
    return slide
