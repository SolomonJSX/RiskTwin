import os
from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from theme import (
    apply_slide_background, add_header, add_footer, add_card, add_icon_badge,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, BORDER_LIGHT, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    # 1. Заголовок слайда с бейджем
    add_header(
        slide,
        tag_text="НЕЙРОСЕТЕВАЯ АРХИТЕКТУРА • DEEP LEARNING",
        title_text="Архитектура ИИ: Двухбашенная нейросеть с кросс-вниманием (PyTorch)",
        subtitle_text="Синтез кредитного портфеля и поведенческой телеметрии для обнаружения стресса за 30–60 дней",
        title_size=Pt(19.5)
    )

    # -------------------------------------------------------------------------
    # ЛЕВАЯ СТОРОНА: Векторная схема архитектуры нейросети (SVG -> PNG)
    # -------------------------------------------------------------------------
    arch_diagram_path = Path(__file__).resolve().parent.parent / "assets" / "dual_tower_architecture.png"
    if arch_diagram_path.exists():
        slide.shapes.add_picture(
            str(arch_diagram_path),
            Inches(0.8), Inches(1.72), Inches(6.9), Inches(5.12)
        )

    # -------------------------------------------------------------------------
    # ПРАВАЯ СТОРОНА: Три технологических преимущества и метрики
    # -------------------------------------------------------------------------

    # 1. Карточка 1: Принцип Двух Башен (Dual-Tower)
    add_card(slide, Inches(7.9), Inches(1.72), Inches(4.65), Inches(1.5), bg_color=RGBColor(238, 242, 255), border_color=RGBColor(199, 210, 254))

    add_icon_badge(
        slide, "database",
        Inches(8.05), Inches(1.88),
        badge_size=Inches(0.52), icon_size=Inches(0.32),
        bg_color=RGBColor(238, 242, 255), border_color=RGBColor(199, 210, 254), color_hex="#4338CA"
    )

    box1 = slide.shapes.add_textbox(Inches(8.72), Inches(1.8), Inches(3.72), Inches(1.35))
    tf1 = box1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

    p1_h = tf1.paragraphs[0]
    p1_h.text = "1. Принцип Двух Башен (Dual-Tower)"
    p1_h.font.name = FONT_MAIN
    p1_h.font.size = Pt(11)
    p1_h.font.bold = True
    p1_h.font.color.rgb = RGBColor(67, 56, 202)

    p1_b = tf1.add_paragraph()
    p1_b.text = (
        "Раздельная токенизация гетерогенных данных: Башня 1 проецирует анкетные и регуляторные параметры "
        "(доход, DTI, ПКБ), а Башня 2 — временные ряды поведения (дни задержки, снятия в ATM, чеки рассрочек). "
        "Каждая башня формирует плотный эмбеддинг d=128."
    )
    p1_b.font.name = FONT_MAIN
    p1_b.font.size = Pt(8.5)
    p1_b.font.color.rgb = COLOR_TEXT_DARK
    p1_b.space_before = Pt(2.5)

    # 2. Карточка 2: Слой Кросс-Внимания (Cross-Attention)
    add_card(slide, Inches(7.9), Inches(3.32), Inches(4.65), Inches(1.5), bg_color=RGBColor(250, 245, 255), border_color=RGBColor(233, 213, 255))

    add_icon_badge(
        slide, "brain",
        Inches(8.05), Inches(3.48),
        badge_size=Inches(0.52), icon_size=Inches(0.32),
        bg_color=RGBColor(250, 245, 255), border_color=RGBColor(233, 213, 255), color_hex="#7E22CE"
    )

    box2 = slide.shapes.add_textbox(Inches(8.72), Inches(3.4), Inches(3.72), Inches(1.35))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0

    p2_h = tf2.paragraphs[0]
    p2_h.text = "2. Механизм Cross-Attention (4 Heads)"
    p2_h.font.name = FONT_MAIN
    p2_h.font.size = Pt(11)
    p2_h.font.bold = True
    p2_h.font.color.rgb = RGBColor(126, 34, 206)

    p2_b = tf2.add_paragraph()
    p2_b.text = (
        "Инновация: Кредитные обязательства заемщика (Query) опрашивают его живое поведение (Key/Value). "
        "Нейросеть находит неочевидную синергию стресса — например, стабильный доход при резком росте "
        "ночных сессий в приложении ДБО (3:00 AM) и хаотичности P2P-переводов."
    )
    p2_b.font.name = FONT_MAIN
    p2_b.font.size = Pt(8.5)
    p2_b.font.color.rgb = COLOR_TEXT_DARK
    p2_b.space_before = Pt(2.5)

    # 3. Карточка 3: Прозрачность и XAI (Explainable AI)
    add_card(slide, Inches(7.9), Inches(4.92), Inches(4.65), Inches(1.42), bg_color=BG_MINT, border_color=BORDER_EMERALD)

    add_icon_badge(
        slide, "shield-check",
        Inches(8.05), Inches(5.08),
        badge_size=Inches(0.52), icon_size=Inches(0.32),
        bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#047857"
    )

    box3 = slide.shapes.add_textbox(Inches(8.72), Inches(5.0), Inches(3.72), Inches(1.25))
    tf3 = box3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_right = tf3.margin_top = tf3.margin_bottom = 0

    p3_h = tf3.paragraphs[0]
    p3_h.text = "3. 100% Объяснимый ИИ (XAI) для Банка"
    p3_h.font.name = FONT_MAIN
    p3_h.font.size = Pt(11)
    p3_h.font.bold = True
    p3_h.font.color.rgb = COLOR_PRIMARY_DARK

    p3_b = tf3.add_paragraph()
    p3_b.text = (
        "Модель не является «черным ящиком». Матрица весов внимания выводит риск-офицеру и аудиторам "
        "АРРФР ключевые драйверы стресса заемщика (рост DTI, падение остатка за 7 дней, просрочки по POS), "
        "соответствуя регуляторным требованиям прозрачности скоринга."
    )
    p3_b.font.name = FONT_MAIN
    p3_b.font.size = Pt(8.4)
    p3_b.font.color.rgb = COLOR_TEXT_DARK
    p3_b.space_before = Pt(2.5)

    # 4. Нижний статус-бар с метриками
    bot_card = add_card(slide, Inches(7.9), Inches(6.42), Inches(4.65), Inches(0.44), bg_color=RGBColor(241, 245, 249), border_color=BORDER_LIGHT)
    bot_box = slide.shapes.add_textbox(Inches(8.0), Inches(6.45), Inches(4.45), Inches(0.38))
    tf_bot = bot_box.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_right = tf_bot.margin_top = tf_bot.margin_bottom = 0
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "⚡ Метрики: ROC-AUC: 0.941 • Инференс: ~7.4 мс • Обучено на 163 517 заемщиках"
    p_bot.font.name = FONT_MAIN
    p_bot.font.size = Pt(8.5)
    p_bot.font.bold = True
    p_bot.font.color.rgb = COLOR_PRIMARY_DARK

    add_footer(slide, current_slide=6, total_slides=12)
    return slide
