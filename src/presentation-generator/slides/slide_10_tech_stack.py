from pptx.util import Inches, Pt
from theme import (
    apply_slide_background, add_header, add_footer, add_card,
    add_icon_badge, add_icon,
    COLOR_PRIMARY, COLOR_PRIMARY_DARK, COLOR_TEXT_DARK, COLOR_TEXT_MUTED,
    BG_CARD, BG_MINT, BORDER_EMERALD, FONT_MAIN
)

def build_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    apply_slide_background(slide)

    add_header(
        slide,
        tag_text="Инженерная зрелость",
        title_text="Технологический стек и готовность к внедрению (TRL 5–6)",
        subtitle_text="Enterprise-архитектура микросервисов, готовая к интеграции в банковский контур и Open Banking РК"
    )

    layers = [
        ("cpu", "AI & Machine Learning Core (Python)",
         "• PyTorch: двухбашенная нейросеть Cross-Attention Net (веса обучены)\n• Google OR-Tools: решатель задач целочисленного линейного программирования (MILP)\n• Scikit-learn & SHAP: конвейер предобработки и объяснимый ИИ (XAI)\n• Микросервисная задержка инференса: < 10 миллисекунд"),
        ("server", "Enterprise Core Backend (C# .NET 10)",
         "• ASP.NET Core 10: принципы Clean Architecture & CQRS\n• Высокопроизводительные шины: gRPC для ML-инференса и SignalR для реал-тайм стриминга\n• База данных: PostgreSQL 16 с оптимизированными индексами и EF Core\n• Регуляторный аудит: хеширование цепочек событий SHA-256 для АРРФР"),
        ("code", "Современный веб и мобильный фронтенд",
         "• Next.js 16 (React 19) с серверным рендерингом (SSR) и мгновенным откликом\n• Tailwind CSS: строгий бело-зеленый банковский UI без визуального шума\n• Recharts & Framer Motion: реактивные радары внимания и графики прогноза ликвидности\n• Готовность к интеграции в мобильные платформы ДБО (iOS / Android)")
    ]

    for i, (icon_name, l_title, l_desc) in enumerate(layers):
        left = Inches(0.8 + i * 4.0)
        card = add_card(slide, left, Inches(1.8), Inches(3.7), Inches(4.0))
        
        # Icon badge
        add_icon_badge(
            slide, icon_name,
            left + Inches(0.2), Inches(1.95),
            badge_size=Inches(0.6), icon_size=Inches(0.34),
            bg_color=BG_MINT, border_color=BORDER_EMERALD, color_hex="#059669"
        )

        box = slide.shapes.add_textbox(left + Inches(0.9), Inches(1.95), Inches(2.6), Inches(0.6))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = l_title
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        desc_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.65), Inches(3.3), Inches(3.0))
        tf_d = desc_box.text_frame
        tf_d.word_wrap = True
        p2 = tf_d.paragraphs[0]
        p2.text = l_desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Banner: Open Banking Readiness with shield-check icon
    add_card(slide, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.8), bg_color=BG_MINT, border_color=BORDER_EMERALD)
    add_icon(slide, "shield-check", Inches(1.0), Inches(6.12), size=Inches(0.35), color_hex="#047857")
    b_box = slide.shapes.add_textbox(Inches(1.45), Inches(6.05), Inches(10.8), Inches(0.6))
    tf_b = b_box.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "Статус зрелости TRL 5–6: Полностью работающий сквозной прототип, проверенный на реальных структурах данных и готовый к пилоту в регуляторной песочнице Нацбанка РК."
    p_b.font.name = FONT_MAIN
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_PRIMARY_DARK

    add_footer(slide, current_slide=10)
    return slide
